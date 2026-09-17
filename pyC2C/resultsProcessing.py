from pptx import Presentation
from pptx.util import Cm
from os import path
import matplotlib.pyplot as plt
from time import strftime
import numpy as np

from .helperFunctions import _str2ms
from .writeTKReport import plotIntScan
from .dictionaries import molecule_mz_cnum


def getPeakEIC(results_df, mz, mz_tolerance, fixed=True):
    """
    Loop through MS1 scans in result data frame and extract intensity of given m/z value.

    Parameters
    ----------
    results_df : pandas.DataFrame
        data frame with combined features and MS1 scans. Must have columns 'MS1' and 'Height'
    mz : float
        m/z value for EIC
    mz_tolerance : float
        m/z tolerance
    fixed : bool
        Use True for fixed tolerance in Da or False for relative tolerance in ppm

    Returns
    -------
    List of relative intensities of MS scans and absolut intensities
    (rel. intensity of MS1 scan multiplied by peak height in result data frame)

    """

    if not fixed:
        mz_tolerance = round(mz * mz_tolerance * 1e-6, ndigits=5)

    # relativ intensity of combined MS1 scan
    eic_rel = [0.0] * len(results_df)

    # relativ intensity multiplied by peak height
    eic_inty = [0.0] * len(results_df)

    # loop through data frame and get MS scans from string
    for i in range(len(results_df)):
        row = results_df.iloc[i, :]

        scan_mz, scan_inty = _str2ms(row['MS1'])

        check_mz = [scan_inty[n] for n in range(len(scan_inty)) if abs(mz - scan_mz[n]) < mz_tolerance]

        if len(check_mz) != 0:
            eic_rel[i] = max(check_mz)
            eic_inty[i] = round(max(check_mz) * row['Height'], ndigits=2)

    return eic_rel, eic_inty


def exportMSScans(results_df, ppt_fout, parameters=None, img_list=None):
    """
    Export all MS scans of peaks in 2D GCxGC data as ppt slides.

    Parameters
    ----------
    results_df : pandas.DataFrame
        data frame with combined features and MS1 scans.
        Must have columns 'Height', 'MS1', 'RT 1d (min)', 'RT 2d (s)', 'Time Cluster ID'
    ppt_fout : str
        file path for ppt file
    parameters : dict (default None)
        dictionary with parameters (ms intensity cut-off, signal picking).
        Set to None, to skip writing parameters to ppt slides.
    img_list : array
        Array with io.Bytes images.
        Set to None, to skip adding images to ppt slides.

    Returns
    -------
    PowerPoint ppt file

    """
    if parameters is not None:

        ms_filepath = parameters['ms filepath']
        sample = parameters['sample']
        inty_threshold = parameters['peak inty threshold']
    else:
        ms_filepath = 'NA'
        sample = 'NA'
        inty_threshold = 0

    presi = Presentation(path.join(path.dirname(__file__), 'template_MSScans.pptx'))

    # title slide
    time_stamp = strftime("_%y%m%d-%H%M%S")

    presi.slides[0].placeholders[17].text = ("%s-%s-%s %s:%s:%s"
                                             % (time_stamp[1:3], time_stamp[3:5], time_stamp[5:7],
                                                time_stamp[8:10], time_stamp[10:12], time_stamp[12:]))

    presi.slides[0].placeholders[14].text = path.basename(ms_filepath)
    presi.slides[0].placeholders[18].text = sample

    # add plots to slides
    if img_list is not None:
        # chromatogram
        presi.slides[1].shapes.add_picture(img_list[0], left=Cm(12.21), top=Cm(2), width=Cm(15.72))
        # signal picking
        presi.slides[2].shapes.add_picture(img_list[1], left=Cm(11.32), top=Cm(3.34), width=Cm(18.73))
        presi.slides[2].shapes.add_picture(img_list[2], left=Cm(11.6), top=Cm(9.98), width=Cm(18.23))

    if parameters is not None:
        presi.slides[1].shapes.placeholders[19].text = str(parameters['scan inty cutoff'])
        presi.slides[1].shapes.placeholders[18].text = str(parameters['modulation time'])
        presi.slides[1].shapes.placeholders[17].text = str(parameters['shift'])
        presi.slides[1].shapes.placeholders[20].text = str(parameters['solvent cut'])

        presi.slides[2].shapes.placeholders[19].text = str(parameters['prominence'])
        presi.slides[2].shapes.placeholders[20].text = str(parameters['minimum peak height'])
        presi.slides[2].shapes.placeholders[21].text = str(parameters['peak inty threshold'])
        presi.slides[2].shapes.placeholders[22].text = str(parameters['modulation time tolerance'])

    # add all MS1 scans with intensity higher than threshold
    k_img = 0

    for i in range(len(results_df)):

        row = results_df.iloc[i, :]

        # check inty of combined peak
        if row['Height'] >= inty_threshold:
            mz, inty = _str2ms(row['MS1'])

            img = plotIntScan(mz, inty, title='cluster id = %i' % row['Time Cluster ID'])
            plt.close()

            t = ("1D: %.2f min\n2D: %.2f s" % (row['RT 1d (min)'], row['RT 2d (s)']))

            if k_img % 4 == 0:
                text_id = 15
                presi.slides.add_slide(presi.slide_layouts[6])
                left = Cm(0)
                top = Cm(0.54)

            elif k_img % 4 == 1:
                text_id = 14
                left = Cm(16.93)
                top = Cm(0.54)

            elif k_img % 4 == 2:
                text_id = 16
                left = Cm(0)
                top = Cm(10.2)
            else:
                text_id = 17
                left = Cm(16.93)
                top = Cm(10.2)

            presi.slides[-1].placeholders[text_id].text = t
            presi.slides[-1].shapes.add_picture(img, left=left, top=top, height=Cm(7.5))

            k_img += 1

    # save ppt
    presi.save(ppt_fout)


def getMoleculeMZ(results_df):

    """
    Loop through MS1 scans and get molecule m/z.

    Parameters
    ----------
    results_df : pandas.DataFrame
        data frame with combined features and MS1 scans. Must have columns 'MS1'

    Returns
    -------
    Array of m/z values of possible molecule peaks

    """

    # array to save m/z values
    molecule_mz = [-1.0] * len(results_df)

    # loop through all integrated MS scans and get last four m/z values
    for i in range(len(results_df)):
        row = results_df.iloc[i, :]

        scan_mz, scan_inty = _str2ms(row['MS1'])

        if len(scan_mz) > 0:

            # get last four signals and check, if the m/z difference is smaller than 3 Da.
            # If so, get m/z with the highest intensity
            candidate_mz = scan_mz[-1]
            candidate_inty = scan_inty[-1]

            for k in [-2, -3, -4]:
                try:
                    if candidate_mz - scan_mz[k] <= 3 and candidate_inty < scan_inty[k]:
                        # overwrite values
                        candidate_mz = scan_mz[k]
                        candidate_inty = scan_inty[k]
                except IndexError:
                    break

            molecule_mz[i] = candidate_mz

    return molecule_mz


def getCNUM_from_MoleculeMZ(molecule_mz, mz_cnum_dict=molecule_mz_cnum):

    cnum = [-1] * len(molecule_mz)

    for i in range(len(molecule_mz)):

        mz = int(molecule_mz[i])

        for k in mz_cnum_dict.keys():
            if mz_cnum_dict[k][0] <= mz <= mz_cnum_dict[k][1]:
                cnum[i] = k
                break

    return cnum


def getDBE(results_df, dbe_lines):
    """
    Calculates the double bound equivalent by first and second retention time of 2D chromatogram.

    Parameters
    ----------
    results_df : pandas.DataFrame
        data frame with combined features and MS1 scans. Must have columns 'RT 1d (min)' and 'RT 2d (s)'
    dbe_lines : list
        List of DBELine objects

    Returns
    -------
    List of DBE
    """

    # empty list for dbe
    dbe = [-1] * len(results_df)

    # loop through result list
    for i in range(len(results_df)):

        temp = []
        diff = []

        rt1 = results_df['RT 1d (min)'].values[i]
        rt2 = results_df['RT 2d (s)'].values[i]

        # loop through list of dbe_lines
        for ln in dbe_lines:

            if abs(ln.linefunc(rt1) - rt2) <= ln.tolerance:
                temp.append(ln.dbe)
                diff.append(abs(ln.linefunc(rt1) - rt2))

        if len(temp) > 1:
            dbe[i] = temp[np.argmin(diff)]
        elif len(temp) == 1:
            dbe[i] = temp[0]

    return dbe


def _checkFragment(mz_list, inty_list, substance_class, fragments_dict, mz_tol=0.2):
    # intensity list of fragments
    frag_inty_list = np.zeros(len(fragments_dict[substance_class]))

    # loop through all given fragments
    for i, frag in enumerate(fragments_dict[substance_class]):

        idx_list = [idx for idx in range(len(mz_list)) if abs(mz_list[idx] - frag) <= mz_tol]

        if len(idx_list) > 0:
            # get maximum intensity of matches
            frag_inty_list[i] = max(inty_list[idx_list])

    return frag_inty_list


def fragments_in_scans(results_df, fragment_dict):

    # create dict for scores
    scores = dict()
    for k in fragment_dict.keys():
        scores[k] = [0.0] * len(results_df)

    # loop through data frame and get scores for all classes in fragment_dict
    for i in range(len(results_df)):

        # get row of data frame
        row = results_df.iloc[i, :]

        # get mz and intensity from ms scan
        scan_mz, scan_inty = _str2ms(row['MS1'])

        for k in scores.keys():
            # get intensity list of matching fragments for every class in dictionary
            frag_inty_list = _checkFragment(scan_mz, scan_inty, k, fragment_dict)

            # use the sum of matching intensity as score
            scores[k][i] = sum(frag_inty_list)

    return scores


