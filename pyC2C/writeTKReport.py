import io
from .helperFunctions import scan_annotation, _str2ms, _dict2str
import matplotlib.pyplot as plt
from os import path
import time

from pptx import Presentation
from pptx.util import Cm


def plotIntScan(mz, inty, title=None):

    fig, ax = plt.subplots(figsize=(5, 3))

    ax.vlines(x=mz, ymin=0, ymax=inty, linewidth = 1)

    idx, max_inty = scan_annotation(mz, inty, threshold=10, min_mz_diff=5)

    for i in idx:
        ax.text(x=mz[i], y=inty[i]+5, s=str(round(mz[i], ndigits=3)), rotation=90, ha='center', fontsize=9)

    ax.set_xlabel('m/z')
    ax.set_ylabel('rel. intensity')
    ax.set_ylim(None, 115)
    ax.spines['top'].set_visible(False)
    ax.spines['right'].set_visible(False)

    if title is not None:
        ax.set_title(title)

    if max(mz) - min(mz) < 5:
        ax.set_xlim(min(mz)-5, max(mz)+5)

    plt.tight_layout()
    img = io.BytesIO()
    plt.savefig(img)

    return img


def init_Report(ms_filepath, sample):

    presi = Presentation(path.join(path.dirname(__file__), 'template_for_reports.pptx'))
    time_stamp = time.strftime("_%y%m%d-%H%M%S")

    presi.slides[0].placeholders[17].text = ("%s-%s-%s %s:%s:%s"
                                             % (time_stamp[1:3], time_stamp[3:5], time_stamp[5:7],
                                                time_stamp[8:10], time_stamp[10:12], time_stamp[12:]))
    presi.slides[0].placeholders[14].text = path.basename(ms_filepath)
    presi.slides[0].placeholders[18].text = sample

    ppt_out = ms_filepath.split('.')[0] + '_report' + time_stamp + '.pptx'
    presi.save(ppt_out)

    return ppt_out


def add_TICSlide(img_io, parameters):

    # open ppt
    presi = Presentation(parameters['ppt filepath'])

    # add tic plot
    presi.slides[1].shapes.add_picture(img_io, left=Cm(12.21), top=Cm(2), width=Cm(15.72))

    # add parameters
    presi.slides[1].shapes.placeholders[19].text = str(parameters['scan inty cutoff'])
    presi.slides[1].shapes.placeholders[18].text = str(parameters['modulation time'])
    presi.slides[1].shapes.placeholders[17].text = str(parameters['shift'])
    presi.slides[1].shapes.placeholders[20].text = str(parameters['solvent cut'])

    presi.save(parameters['ppt filepath'])


def add_PeakfinderSlide(img_peakfinder, img_inty, img_cluster, parameters):

    # open ppt
    presi = Presentation(parameters['ppt filepath'])

    # add plots
    presi.slides[2].shapes.add_picture(img_peakfinder, left=Cm(11.32), top=Cm(3.34), width=Cm(18.73))
    presi.slides[2].shapes.add_picture(img_inty, left=Cm(11.6), top=Cm(9.98), width=Cm(18.23))

    # add parameters
    presi.slides[2].shapes.placeholders[19].text = str(parameters['prominence'])
    presi.slides[2].shapes.placeholders[20].text = str(parameters['minimum peak height'])
    presi.slides[2].shapes.placeholders[21].text = str(parameters['peak inty threshold'])
    presi.slides[2].shapes.placeholders[22].text = str(parameters['modulation time tolerance'])

    # add plots with cluster ids
    presi.slides[3].shapes.add_picture(img_cluster, left=Cm(0.89), top=Cm(1.33), width=Cm(31.01))

    presi.save(parameters['ppt filepath'])


def add_DBESlide(img_lines, img_bar, img_dbe, dbe_param_string, parameters):

    # open ppt
    presi = Presentation(parameters['ppt filepath'])

    # add plots
    presi.slides[4].shapes.add_picture(img_lines, left=Cm(10.84), top=Cm(3.11), width=Cm(22.78))

    # add parameters of dbe lines
    presi.slides[4].shapes.placeholders[18].text = dbe_param_string

    # add plots
    presi.slides[5].shapes.add_picture(img_bar, left=Cm(0), top=Cm(6.39), width=Cm(10.19))
    presi.slides[5].shapes.add_picture(img_dbe, left=Cm(11.09), top=Cm(3.03), width=Cm(22.78))

    presi.save(parameters['ppt filepath'])


def add_listSlide(img_class, data_results, parameters):

    presi = Presentation(parameters['ppt filepath'])

    # fragment list
    presi.slides[6].shapes.placeholders[18].text = _dict2str(data_results.fragments_dict)

    # lri list
    presi.slides[6].shapes.placeholders[19].text = _dict2str(data_results.lri_dict, with_tabs=True)

    # add class plot
    presi.slides[7].shapes.add_picture(img_class, left=Cm(0.28), top=Cm(1.33), width=Cm(30.98))

    presi.save(parameters['ppt filepath'])


def add_scans(data_results, parameters):

    presi = Presentation(parameters['ppt filepath'])
    k_img = 0

    for i in range(len(data_results.results)):

        row = data_results.results.iloc[i, :]

        # check inty of combined peak
        if row['Height'] >= parameters['peak inty threshold']:

            mz, inty = _str2ms(row['MS1'])

            img = plotIntScan(mz, inty, title='cluster id = %i' % row['Time Cluster ID'])
            plt.close()

            t = ("1D: %.2f min\n2D: %.2f s\nLIR: %i \nDBE: %i"
                 "\nPI: %.1f\nO: %.1f\nN: %.1f\nA: %.1f") % (row['RT 1d (min)'],
                                                             row['RT 2d (s)'],
                                                             row['LRI'],
                                                             row['DBE'],
                                                             row['Score PI'],
                                                             row['Score O'],
                                                             row['Score N'],
                                                             row['Score A'])

            if k_img % 4 == 0:
                text_id = 15
                presi.slides.add_slide(presi.slide_layouts[6])
                left = Cm(0)
                top = Cm(0.54)

            elif k_img % 4 == 1:
                text_id = 14
                left = Cm(16.93)
                top = Cm(0.54)

            elif k_img % 4 == 2:
                text_id = 16
                left = Cm(0)
                top = Cm(10.2)
            else:
                text_id = 17
                left = Cm(16.93)
                top = Cm(10.2)

            presi.slides[-1].placeholders[text_id].text = t
            presi.slides[-1].shapes.add_picture(img, left=left, top=top, height=Cm(7.5))

            k_img += 1

    # overwrite
    presi.save(parameters['ppt filepath'])




